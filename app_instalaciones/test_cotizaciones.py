from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
import shutil
import tempfile
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from app_instalaciones.models import (
    CotizacionServicioTecnico,
    Ejecutivo,
    Mantenimiento,
    Tecnico,
)
from app_instalaciones.admin import CotizacionServicioTecnicoAdminForm


class SeguimientoCotizacionesTests(TestCase):
    def setUp(self):
        self.media_root = tempfile.mkdtemp(prefix="sesur-cotizaciones-", dir=settings.BASE_DIR)
        self.media_override = override_settings(MEDIA_ROOT=self.media_root)
        self.media_override.enable()
        self.addCleanup(self.media_override.disable)
        self.addCleanup(shutil.rmtree, self.media_root, True)
        tecnico_user = User.objects.create_user("tecnico_cotizacion")
        tecnico = Tecnico.objects.create(nombre="Técnico cotización", usuario=tecnico_user)
        servicio = Mantenimiento.objects.create(
            codigo="CLI-10", numero_ticket="TKT-20", cliente="Cliente",
            direccion="Dirección", tecnico=tecnico,
        )
        self.cotizacion = CotizacionServicioTecnico.objects.create(
            mantenimiento=servicio, tecnico=tecnico, codigo="CLI-10",
            registro="TKT-20", cliente="Cliente", direccion="Dirección",
            descripcion="Equipos requeridos", soporte="cotizaciones/soporte.pdf",
            creado_por=tecnico_user,
        )
        self.ejecutivo = Ejecutivo.objects.create(nombre="Ejecutivo comercial")
        admin = User.objects.create_user("admin_cotizacion")
        grupo, _ = Group.objects.get_or_create(name="Cotizaciones")
        admin.groups.add(grupo)
        self.client.force_login(admin)

    def test_modulo_restringido_al_rol_cotizaciones(self):
        usuario_sin_permiso = User.objects.create_user("sin_permiso_cotizaciones")
        grupo, _ = Group.objects.get_or_create(name="Administrador")
        usuario_sin_permiso.groups.add(grupo)
        self.client.force_login(usuario_sin_permiso)

        inicio = self.client.get(reverse("home"))
        listado = self.client.get(reverse("listar_cotizaciones_tecnicas"))

        self.assertNotContains(inicio, "Ver cotizaciones")
        self.assertRedirects(
            listado,
            f'{reverse("home")}?next={reverse("listar_cotizaciones_tecnicas")}',
        )

    def test_muestra_alerta_emergente_por_orden_nueva(self):
        inicio = self.client.get(reverse("home"))

        self.assertContains(inicio, "Nueva orden para cotización")
        self.assertContains(inicio, "Se asignó una nueva orden para cotización.")
        self.assertContains(inicio, "Revisar cotizaciones")

    def test_envio_seguimiento_y_respuesta(self):
        self.client.post(
            reverse("gestionar_cotizacion_tecnica", args=[self.cotizacion.pk]),
            {"accion": "ENVIAR", "ejecutivo": self.ejecutivo.pk,
             "medio_envio": "CORREO", "dirigida_a": "Compras",
             "numero_cotizacion": "COT-2026-001"},
        )
        self.cotizacion.refresh_from_db()
        self.assertEqual(self.cotizacion.estado, "PENDIENTE")
        self.assertEqual(self.cotizacion.numero_cotizacion, "COT-2026-001")
        self.assertIsNotNone(self.cotizacion.fecha_envio)

        self.cotizacion.fecha_orden = timezone.now() - timedelta(days=8)
        self.cotizacion.save(update_fields=["fecha_orden"])
        listado = self.client.get(reverse("listar_cotizaciones_tecnicas"))
        self.assertContains(listado, "Tienes 1 cotización enviada sin respuesta")
        self.assertContains(listado, "Primer seguimiento")
        inicio = self.client.get(reverse("home"))
        self.assertContains(inicio, "Tienes 1 cotización enviada hace 8 días")

        self.client.post(
            reverse("gestionar_cotizacion_tecnica", args=[self.cotizacion.pk]),
            {"accion": "GESTIONADA", "nota_seguimiento": "Se llamó al cliente."},
        )
        self.cotizacion.refresh_from_db()
        self.assertEqual(self.cotizacion.estado, "PENDIENTE")
        seguimiento = self.cotizacion.seguimientos.get()
        self.assertEqual(seguimiento.numero, 1)
        self.assertEqual(seguimiento.nota, "Se llamó al cliente.")
        self.client.post(
            reverse("gestionar_cotizacion_tecnica", args=[self.cotizacion.pk]),
            {"accion": "APROBAR"},
        )
        self.cotizacion.refresh_from_db()
        self.assertEqual(self.cotizacion.estado, "APROBADA")
        self.assertEqual(self.cotizacion.accion_seguimiento, "GESTIONADA")
        self.assertIsNotNone(self.cotizacion.fecha_gestion)

    def test_rechazo_exige_nota(self):
        self.cotizacion.estado = "PENDIENTE"
        self.cotizacion.fecha_envio = timezone.now()
        self.cotizacion.save(update_fields=["estado", "fecha_envio"])
        self.client.post(
            reverse("gestionar_cotizacion_tecnica", args=[self.cotizacion.pk]),
            {"accion": "RECHAZAR", "nota_rechazo": ""},
        )
        self.cotizacion.refresh_from_db()
        self.assertEqual(self.cotizacion.estado, "PENDIENTE")

        self.client.post(
            reverse("gestionar_cotizacion_tecnica", args=[self.cotizacion.pk]),
            {"accion": "RECHAZAR", "nota_rechazo": "Fuera del presupuesto."},
        )
        self.cotizacion.refresh_from_db()
        self.assertEqual(self.cotizacion.estado, "RECHAZADA")
        self.assertEqual(self.cotizacion.nota_rechazo, "Fuera del presupuesto.")

    def test_creacion_manual_y_contadores_como_filtros(self):
        respuesta = self.client.post(
            reverse("crear_cotizacion_tecnica_manual"),
            {
                "codigo": "MAN-01",
                "registro": "REG-MANUAL",
                "fecha_orden": "2026-10-08T10:02",
                "cliente": "Cliente manual",
                "direccion": "Calle manual",
                "contacto": "Persona de compras",
                "telefono": "3001234567",
                "descripcion": "Cotización creada desde plataforma",
                "soporte": SimpleUploadedFile(
                    "manual.pdf", b"%PDF-1.4\nprueba", content_type="application/pdf"
                ),
            },
        )
        self.assertEqual(respuesta.status_code, 302)
        manual = CotizacionServicioTecnico.objects.get(codigo="MAN-01")
        self.assertIsNone(manual.mantenimiento)
        self.assertEqual(manual.registro, "REG-MANUAL")
        self.assertEqual(manual.dirigida_a, "Persona de compras")
        self.assertEqual(manual.telefono, "3001234567")
        self.assertEqual(manual.fecha_orden.date().isoformat(), "2026-10-08")

        self.cotizacion.estado = "APROBADA"
        self.cotizacion.fecha_envio = timezone.now()
        self.cotizacion.save(update_fields=["estado", "fecha_envio"])
        aprobadas = self.client.get(
            reverse("listar_cotizaciones_tecnicas"), {"estado": "APROBADA"}
        )
        self.assertContains(aprobadas, "CLI-10")
        self.assertNotContains(aprobadas, "MAN-01")
        self.assertContains(aprobadas, ">1</span>", html=False)

    def test_no_permite_repetir_el_numero_de_registro_manual(self):
        respuesta = self.client.post(
            reverse("crear_cotizacion_tecnica_manual"),
            {
                "codigo": "MAN-02",
                "registro": self.cotizacion.registro,
                "fecha_orden": "2026-10-08T10:02",
                "cliente": "Cliente duplicado",
                "direccion": "Calle duplicada",
                "descripcion": "No debe crearse",
                "soporte": SimpleUploadedFile(
                    "duplicado.pdf", b"%PDF-1.4\n", content_type="application/pdf"
                ),
            },
            follow=True,
        )

        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, "No se pudo procesar la orden")
        self.assertContains(
            respuesta,
            f"El registro {self.cotizacion.registro} ya tiene una cotización.",
        )
        self.assertContains(respuesta, 'formularioConError = "manual"')
        self.assertEqual(
            CotizacionServicioTecnico.objects.filter(
                registro=self.cotizacion.registro
            ).count(),
            1,
        )

    @patch(
        "app_instalaciones.home_views.user_views.analizar_visita_tecnica_cotizacion"
    )
    def test_importa_orden_pdf_y_conserva_soporte(self, analizar):
        analizar.return_value = {
            "registro": "1107470",
            "fecha_orden": timezone.make_aware(
                datetime(2026, 10, 8, 10, 2)
            ),
            "tecnico_nombre": "Técnico cotización",
            "codigo": "3480",
            "cliente": "EDIFICIO LEFACOP",
            "direccion": "CARRERA 29A 18 - 09",
            "descripcion": "50 m de cable UTP y 15 canaletas",
            "actividad_realizada": "Visita para reubicación de sensores.",
            "telefono": "3146313992",
            "contacto": "Juanita Portilla",
            "omt": "SARITA ORTEGA",
        }
        respuesta = self.client.post(
            reverse("importar_cotizacion_visita_pdf"),
            {"soporte": SimpleUploadedFile(
                "visita.pdf", b"%PDF-1.4\norden", content_type="application/pdf"
            )},
        )
        self.assertEqual(respuesta.status_code, 302)
        importada = CotizacionServicioTecnico.objects.get(registro="1107470")
        self.assertEqual(importada.codigo, "3480")
        self.assertEqual(importada.cliente, "EDIFICIO LEFACOP")
        self.assertEqual(importada.contacto, "Juanita Portilla")
        self.assertEqual(importada.dirigida_a, "Juanita Portilla")
        self.assertEqual(importada.telefono, "3146313992")
        self.assertIsNotNone(importada.tecnico)
        self.assertEqual(importada.fecha_orden.date().isoformat(), "2026-10-08")
        self.assertTrue(importada.soporte.name.endswith("visita.pdf"))
        self.assertEqual(bytes(importada.soporte_contenido), b"%PDF-1.4\norden")

        descarga = self.client.get(
            reverse("descargar_soporte_cotizacion", args=[importada.pk])
        )
        self.assertEqual(descarga.status_code, 200)
        self.assertEqual(descarga["Content-Type"], "application/pdf")
        self.assertEqual(b"".join(descarga.streaming_content), b"%PDF-1.4\norden")

    @patch(
        "app_instalaciones.home_views.user_views.analizar_visita_tecnica_cotizacion"
    )
    def test_importa_orden_sin_codigo_usando_cero(self, analizar):
        analizar.return_value = {
            "registro": "1107471",
            "fecha_orden": timezone.now(),
            "tecnico_nombre": "Técnico cotización",
            "codigo": "",
            "cliente": "CLIENTE SIN CÓDIGO",
            "direccion": "CALLE 1",
            "descripcion": "Cotización requerida",
            "actividad_realizada": "Visita técnica",
            "telefono": "",
            "contacto": "",
            "omt": "",
        }
        respuesta = self.client.post(
            reverse("importar_cotizacion_visita_pdf"),
            {"soporte": SimpleUploadedFile(
                "sin-codigo.pdf", b"%PDF-1.4\norden", content_type="application/pdf"
            )},
        )

        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(
            CotizacionServicioTecnico.objects.get(registro="1107471").codigo,
            "0",
        )

    def test_busca_cliente_y_direccion_por_codigo_de_mantenimiento(self):
        respuesta = self.client.get(
            reverse("buscar_cliente_mantenimiento"), {"codigo": "CLI-10"}
        )
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(respuesta.json()["encontrado"])
        self.assertEqual(respuesta.json()["cliente"], "Cliente")
        self.assertEqual(respuesta.json()["direccion"], "Dirección")

    def test_mantenimiento_no_es_obligatorio_en_admin(self):
        formulario = CotizacionServicioTecnicoAdminForm()
        self.assertFalse(formulario.fields["mantenimiento"].required)

    def test_soporte_antiguo_faltante_no_genera_error_500(self):
        respuesta = self.client.get(
            reverse("descargar_soporte_cotizacion", args=[self.cotizacion.pk]),
            follow=True,
        )

        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, "El soporte PDF original no está disponible")
        self.assertContains(respuesta, "no se puede reemplazar")

    def test_soporte_antiguo_existente_se_copia_a_base_de_datos(self):
        contenido = b"%PDF-1.4\nsoporte original"
        self.cotizacion.soporte.save(
            "soporte-original.pdf",
            SimpleUploadedFile("soporte-original.pdf", contenido),
        )
        self.cotizacion.save(update_fields=["soporte"])

        respuesta = self.client.get(
            reverse("descargar_soporte_cotizacion", args=[self.cotizacion.pk])
        )

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(b"".join(respuesta.streaming_content), contenido)
        self.cotizacion.refresh_from_db()
        self.assertEqual(bytes(self.cotizacion.soporte_contenido), contenido)
        self.assertEqual(self.cotizacion.soporte_nombre, "soporte-original.pdf")

    def test_recupera_soporte_antiguo_por_registro_si_la_ruta_cambio(self):
        contenido = b"%PDF-1.4\nsoporte original por registro"
        ruta = (
            Path(self.media_root)
            / "cotizaciones_tecnicas"
            / "2026"
            / "10"
            / "VISITA_TECNICA_TKT-20_CLIENTE.pdf"
        )
        ruta.parent.mkdir(parents=True)
        ruta.write_bytes(contenido)

        respuesta = self.client.get(
            reverse("descargar_soporte_cotizacion", args=[self.cotizacion.pk])
        )

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(b"".join(respuesta.streaming_content), contenido)
        self.cotizacion.refresh_from_db()
        self.assertEqual(bytes(self.cotizacion.soporte_contenido), contenido)

    def test_exporta_cotizaciones_en_excel(self):
        respuesta = self.client.get(reverse("exportar_cotizaciones_excel"))

        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("spreadsheetml", respuesta["Content-Type"])
        libro = load_workbook(BytesIO(respuesta.content), read_only=True)
        hoja = libro["Cotizaciones"]
        self.assertEqual(hoja.cell(1, 3).value, "Registro / solicitante")
        self.assertEqual(hoja.cell(2, 3).value, "TKT-20")

