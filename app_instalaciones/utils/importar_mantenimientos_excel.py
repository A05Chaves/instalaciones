import re
import unicodedata
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app_instalaciones.models import CuadroInsta, Mantenimiento, Tecnico
from app_instalaciones.utils.importar_mantenimientos_pdf import (
    FALLAS_MAP, SERVICIOS_MAP, limpiar_texto, normalizar_clave,
)


def _normalizar_encabezado(valor):
    texto = limpiar_texto(str(valor or "")).upper()
    texto = "".join(c for c in unicodedata.normalize("NFD", texto)
                    if unicodedata.category(c) != "Mn")
    return re.sub(r"[^A-Z0-9]+", " ", texto).strip()


ALIAS_COLUMNAS = {
    "numero_ticket": {"TICKET", "NUMERO TICKET", "N TICKET", "NO TICKET", "NUMERO DE TICKET"},
    "estado_ticket": {"ESTADO", "ESTADO TICKET", "ESTADO DEL TICKET"},
    "fecha_creacion_servicio": {"FECHA REGISTRO", "FECHA CREADO", "FECHA DE CREADO", "CREADO", "FECHA CREACION"},
    "codigo": {"CODIGO", "CODIGO CLIENTE", "CODIGO DEL CLIENTE"},
    "cliente": {"CLIENTE", "NOMBRE CLIENTE"},
    "ciudad": {"CIUDAD", "MUNICIPIO"},
    "direccion": {"DIRECCION", "DIRECCION CLIENTE"},
    "tipo_servicio": {"TIPO SERVICIO", "SERVICIO", "TIPO DE SERVICIO"},
    "tipo_falla": {"TIPO FALLA", "FALLA", "TRABAJO", "TIPO DE FALLA"},
    "tecnico_nombre": {"TECNICO", "NOMBRE TECNICO"},
    "fecha_inicio": {"INICIO", "FECHA INICIO", "INICIO TECNICO"},
    "fecha_fin": {"FIN", "FECHA FIN", "FIN TECNICO"},
    "hora_entrada": {"HORA ENTRADA", "ENTRADA"},
    "hora_salida": {"HORA SALIDA", "SALIDA"},
    "horas": {"DURACION", "HORAS", "TIEMPO"},
    "codigo_acta": {"ORDEN", "CODIGO ACTA", "NUMERO ORDEN", "NUMERO DE ORDEN"},
    "realizado": {"REALIZADO", "FECHA REALIZADO", "FECHA DE REALIZACION"},
    "novedad": {"NOVEDAD", "NOVEDAD TECNICO"},
    "observacion": {"OBSERVACION", "OBSERVACIONES", "DESCRIPCION"},
    "pendiente": {"PENDIENTE", "PENDIENTES"},
    "omt": {"OMT"},
}


def _mapa_columnas(fila):
    inversos = {alias: campo for campo, aliases in ALIAS_COLUMNAS.items() for alias in aliases}
    return {inversos[encabezado]: i for i, valor in enumerate(fila)
            if (encabezado := _normalizar_encabezado(valor)) in inversos}


def _a_datetime(valor):
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime.combine(valor, time.min)
    texto = limpiar_texto(str(valor))
    for formato in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
                    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return None


def _a_hora(valor):
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.time().replace(microsecond=0)
    if isinstance(valor, time):
        return valor.replace(microsecond=0)
    for formato in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(limpiar_texto(str(valor)), formato).time()
        except ValueError:
            continue
    return None


def _a_horas(valor):
    if valor in (None, "", "---"):
        return None
    if isinstance(valor, time):
        return Decimal(valor.hour) + Decimal(valor.minute) / 60 + Decimal(valor.second) / 3600
    texto = limpiar_texto(str(valor))
    try:
        if ":" in texto:
            partes = [Decimal(parte) for parte in texto.split(":")]
            return partes[0] + partes[1] / 60 + (partes[2] / 3600 if len(partes) > 2 else 0)
        return Decimal(texto.replace(",", "."))
    except (InvalidOperation, IndexError):
        return None


def _valor(fila, columnas, campo):
    indice = columnas.get(campo)
    return fila[indice] if indice is not None and indice < len(fila) else None


def _texto(fila, columnas, campo):
    valor = _valor(fila, columnas, campo)
    return limpiar_texto(str(valor)) if valor is not None else ""


def analizar_excel_mantenimientos(archivo_excel, actualizar_existentes=False):
    try:
        libro = load_workbook(archivo_excel, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError) as error:
        raise ValueError("El archivo Excel está dañado o no tiene formato .xlsx válido.") from error
    try:
        filas = list(libro.active.iter_rows(values_only=True))
        indice_encabezado = None
        columnas = {}
        for indice, fila in enumerate(filas[:15]):
            candidatas = _mapa_columnas(fila)
            if "numero_ticket" in candidatas and len(candidatas) >= 3:
                indice_encabezado, columnas = indice, candidatas
                break
        if indice_encabezado is None:
            raise ValueError("No se encontró una columna de número de ticket en el archivo Excel.")

        tecnicos = {normalizar_clave(t.nombre).replace(" ", ""): t for t in Tecnico.objects.all()}
        resultados = []
        for numero_fila, fila in enumerate(filas[indice_encabezado + 1:], start=indice_encabezado + 2):
            ticket = _texto(fila, columnas, "numero_ticket")
            if not ticket and not any(valor not in (None, "") for valor in fila):
                continue
            servicio_original = _texto(fila, columnas, "tipo_servicio")
            falla_original = _texto(fila, columnas, "tipo_falla")
            servicios_choices = {
                normalizar_clave(etiqueta): valor
                for valor, etiqueta in Mantenimiento.TIPO_SERVICIO_CHOICES
            }
            fallas_choices = {
                normalizar_clave(etiqueta).replace(".", ""): valor
                for valor, etiqueta in Mantenimiento.TIPO_FALLA_CHOICES
            }
            servicio_clave = normalizar_clave(servicio_original)
            falla_clave = normalizar_clave(falla_original)
            servicio = servicios_choices.get(
                servicio_clave, SERVICIOS_MAP.get(servicio_clave, servicio_clave)
            )
            falla = fallas_choices.get(
                falla_clave.replace(".", ""), FALLAS_MAP.get(falla_clave, falla_clave)
            )
            if servicio not in dict(Mantenimiento.TIPO_SERVICIO_CHOICES):
                servicio = "MANTENIMIENTO CORRECTIVO"
            if falla not in dict(Mantenimiento.TIPO_FALLA_CHOICES):
                falla = "OTRO" if falla_original else ""

            inicio = _a_datetime(_valor(fila, columnas, "fecha_inicio"))
            fin = _a_datetime(_valor(fila, columnas, "fecha_fin"))
            tecnico_nombre = _texto(fila, columnas, "tecnico_nombre")
            tecnico = tecnicos.get(normalizar_clave(tecnico_nombre).replace(" ", "")) if tecnico_nombre else None
            codigo = _texto(fila, columnas, "codigo")
            instalacion = CuadroInsta.objects.filter(codigo__iexact=codigo).first() if codigo else None
            existente = bool(ticket and Mantenimiento.objects.filter(numero_ticket=ticket).exists())
            fecha_realizado = _a_datetime(_valor(fila, columnas, "realizado"))
            data = {
                "numero_ticket": ticket, "estado_ticket": _texto(fila, columnas, "estado_ticket"),
                "codigo": codigo,
                "cliente": _texto(fila, columnas, "cliente") or (instalacion.cliente if instalacion else ""),
                "ciudad": _texto(fila, columnas, "ciudad") or (instalacion.ciudad if instalacion else ""),
                "direccion": _texto(fila, columnas, "direccion") or (instalacion.direccion if instalacion else ""),
                "tipo_servicio": servicio, "tipo_falla": falla,
                "tecnico_nombre": tecnico_nombre, "tecnico_id": tecnico.id if tecnico else None,
                "fecha_creacion_servicio": _a_datetime(_valor(fila, columnas, "fecha_creacion_servicio")),
                "fecha_inicio": inicio, "fecha_fin": fin,
                "hora_entrada": _a_hora(_valor(fila, columnas, "hora_entrada")) or (inicio.time() if inicio else None),
                "hora_salida": _a_hora(_valor(fila, columnas, "hora_salida")) or (fin.time() if fin else None),
                "horas": _a_horas(_valor(fila, columnas, "horas")),
                "realizado": fecha_realizado.date() if fecha_realizado else (fin.date() if fin else None),
                "codigo_acta": _texto(fila, columnas, "codigo_acta"),
                "problema_solucionado": "", "cotizacion": "",
                "novedad": _texto(fila, columnas, "novedad"),
                "observacion": _texto(fila, columnas, "observacion"),
                "pendiente": _texto(fila, columnas, "pendiente"), "omt": _texto(fila, columnas, "omt"),
            }
            advertencias = []
            if not ticket:
                advertencias.append("No se encontró número de ticket.")
            if existente:
                advertencias.append("Ticket existente: se actualizará." if actualizar_existentes else "Ticket ya existe en la base de datos.")
            if tecnico_nombre and not tecnico:
                advertencias.append(f"Se creará el técnico: {tecnico_nombre}")
            elif not tecnico_nombre:
                advertencias.append("No se encontró técnico.")
            if not codigo:
                advertencias.append("No se encontró código.")
            resultados.append({"pagina": numero_fila, "data": data, "advertencias": advertencias,
                               "importable": bool(ticket) and (actualizar_existentes or not existente)})
        return resultados
    finally:
        libro.close()
