from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from flask import request
from app_instalaciones.home_views.user_views import (
    puede_editar_instalacion, puede_programar
)
from app_instalaciones.models.forms import CuadroInstaForm, ExcelUploadForm
from app_instalaciones.models.cuadroInstalaciones import CuadroInsta, RegistroImportacion, Tecnico, Ejecutivo
from django.shortcuts import get_object_or_404
import openpyxl
import csv
from django.http import HttpResponse
import pandas as pd
from datetime import datetime, date, time
from django.utils import timezone
from app_instalaciones.models.cuadroInstalaciones import CuadroInsta, Ciudad
from django.db.models import Q

# PAGINA INICIAL DEL PROYECTO
# VERSION 4 PARA EDICION SESUR 28 DE ABRIL 2025


def home(request):
    if not request.user.is_authenticated:
        return redirect('login')  # redirige al login si no está autenticado
    # muestra home si ya está autenticado
    return render(request, "home.html")


@login_required(login_url='login')
@user_passes_test(puede_programar)
def registro_inst(request):
    if request.user.groups.filter(name='Visor').exists():
        messages.warning(
            request, " No tienes permisos para registrar instalaciones.")
        return redirect('home')

    if request.method == 'POST':
        form = CuadroInstaForm(request.POST)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.usuario = request.user

            # Asignar FKs explícitamente por si el template tenía campos disabled
            t1 = request.POST.get('tecnico1') or None
            t2 = request.POST.get('tecnico2') or None
            ej = request.POST.get('ejecutivo') or None

            # Con *_id asignas por PK sin tener que hacer .get()
            obj.tecnico1_id = t1
            obj.tecnico2_id = t2
            obj.ejecutivo_id = ej

            obj.save()
            messages.success(request, 'Instalación registrada con éxito.')
            return redirect('registro_inst')
        else:
            messages.error(
                request, 'Por favor, corrige los errores en el formulario.')
    else:
        form = CuadroInstaForm()

    return render(request, 'registro_inst.html', {'form': form})

# VISTA PARA LISTAR INSTALACIONES


def lista_instalaciones(request):
    ejecutivos_seleccionados = request.GET.getlist('ejecutivo')
    tecnicos_seleccionados = request.GET.getlist('tecnico')
    bodega = request.GET.get('bodega')
    alistado = request.GET.get('alistado')
    facturado = request.GET.get('facturado')
    ciudades_seleccionadas = request.GET.getlist('ciudad')
    estados_seleccionados = request.GET.getlist('estado')
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    q = request.GET.get('q', '').strip()

    qs = (
        CuadroInsta.objects
        .select_related('tecnico1', 'tecnico2', 'ejecutivo', 'usuario')
        .order_by('-id')
    )

    if fecha_inicio:
        qs = qs.filter(fecha__date__gte=fecha_inicio)

    if fecha_fin:
        qs = qs.filter(fecha__date__lte=fecha_fin)

    if q:
        terminos = [t.strip() for t in q.split(',') if t.strip()]
        filtro_total = Q()

        for termino in terminos:
            filtro_total &= (
                Q(pvg__icontains=termino) |
                Q(fecha__icontains=termino) |
                Q(codigo__icontains=termino) |
                Q(cliente__icontains=termino) |
                Q(ciudad__icontains=termino) |
                Q(direccion__icontains=termino) |
                Q(instalacion__icontains=termino) |
                Q(dias_cotizados__icontains=termino) |
                Q(en_bodega__icontains=termino) |
                Q(fecha_inicio__icontains=termino) |
                Q(fecha_terminacion__icontains=termino) |
                Q(finaliza__icontains=termino) |
                Q(orden__icontains=termino) |
                Q(tecnico1__nombre__icontains=termino) |
                Q(tecnico2__nombre__icontains=termino) |
                Q(estado__icontains=termino) |
                Q(observacion__icontains=termino) |
                Q(ejecutivo__nombre__icontains=termino) |
                Q(usuario__username__icontains=termino) |
                Q(updated_at__icontains=termino) |
                Q(fecha_alistado__icontains=termino) |
                Q(fecha_facturacion__icontains=termino)
            )

        qs = qs.filter(filtro_total)

    if ciudades_seleccionadas:
        qs = qs.filter(ciudad__in=ciudades_seleccionadas)

    if estados_seleccionados:
        qs = qs.filter(estado__in=estados_seleccionados)

    if ejecutivos_seleccionados:
        qs = qs.filter(ejecutivo_id__in=ejecutivos_seleccionados)

    if tecnicos_seleccionados:
        qs = qs.filter(
            Q(tecnico1_id__in=tecnicos_seleccionados) |
            Q(tecnico2_id__in=tecnicos_seleccionados)
        )

    if bodega:
        qs = qs.filter(en_bodega=bodega)

    if alistado:
        qs = qs.filter(alistado=(alistado == "SI"))

    if facturado:
        qs = qs.filter(facturado=(facturado == "SI"))

    if (
        not fecha_inicio
        and not fecha_fin
        and not q
        and not ciudades_seleccionadas
        and not estados_seleccionados
        and not ejecutivos_seleccionados
        and not tecnicos_seleccionados
        and not bodega
        and not alistado
        and not facturado
    ):
        qs = qs[:200]

    instalaciones = list(qs)
    total_pvg = CuadroInsta.objects.count()

    ciudades = (
        CuadroInsta.objects
        .exclude(ciudad__isnull=True)
        .exclude(ciudad="")
        .values_list('ciudad', flat=True)
        .distinct()
        .order_by('ciudad')
    )

    estados = (
        CuadroInsta.objects
        .exclude(estado__isnull=True)
        .exclude(estado="")
        .values_list('estado', flat=True)
        .distinct()
        .order_by('estado')
    )

    ejecutivos = Ejecutivo.objects.all().order_by('nombre')
    tecnicos = Tecnico.objects.all().order_by('nombre')

    return render(
        request,
        'lista_instalaciones.html',
        {
            'instalaciones': instalaciones,
            'fecha_inicio': fecha_inicio,
            'fecha_fin': fecha_fin,
            'q': q,
            'total_pvg': total_pvg,
            'ciudades': ciudades,
            'estados': estados,
            'ciudades_seleccionadas': ciudades_seleccionadas,
            'estados_seleccionados': estados_seleccionados,
            'ejecutivos': ejecutivos,
            'tecnicos': tecnicos,
            'ejecutivos_seleccionados': [int(x) for x in ejecutivos_seleccionados],
            'tecnicos_seleccionados': [int(x) for x in tecnicos_seleccionados],
            'bodega': bodega,
            'alistado': alistado,
            'facturado': facturado,
        }
    )

# AGREGADO 25 DE ABRIL


def is_editor_or_admin(user):
    """ Verifica si el usuario es administrador o pertenece al grupo de editores """
    return user.is_staff or user.groups.filter(name='editor').exists()


@login_required
@user_passes_test(puede_editar_instalacion)
def editar_instalacion(request, id):
    instalacion = get_object_or_404(CuadroInsta, pk=id)

    # Bloquea si está anulada (salvo superusuario)
    if instalacion.estado and instalacion.estado.strip().upper() == "ANULADO" and not request.user.is_superuser:
        messages.warning(
            request, "No puedes editar una instalación que ha sido anulada.")
        return redirect(
            request.META.get(
                'HTTP_REFERER',
                'lista_instalaciones'
            )
        )

    if request.method == 'POST':
        form = CuadroInstaForm(
            request.POST, instance=instalacion, user=request.user)
        if form.is_valid():
            try:
                obj = form.save(commit=False)

                #  Asegurar FKs aunque los selects estén disabled en el form
                t1 = request.POST.get('tecnico1')
                t2 = request.POST.get('tecnico2')
                ej = request.POST.get('ejecutivo')

                # Si no viene en POST (por algún motivo), conservar el valor anterior
                obj.tecnico1_id = t1 if t1 not in (None, "") else getattr(
                    instalacion, "tecnico1_id", None)
                obj.tecnico2_id = t2 if t2 not in (None, "") else getattr(
                    instalacion, "tecnico2_id", None)
                obj.ejecutivo_id = ej if ej not in (None, "") else getattr(
                    instalacion, "ejecutivo_id", None)

                obj.save()

                messages.success(
                    request, "Instalación actualizada correctamente.")

                next_url = request.POST.get('next') or request.GET.get('next')

                if next_url:
                    return redirect(next_url)

                return redirect('lista_instalaciones')
            except Exception as e:
                messages.error(request, f"Error al guardar: {str(e)}")
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"Error en {field}: {error}")
    else:
        form = CuadroInstaForm(instance=instalacion, user=request.user)

    return render(request, 'editar_instalacion.html', {
        'form': form,
        'instalacion': instalacion,
        'next': request.GET.get('next', ''),
    })


@login_required
def eliminar_instalacion(request, id):
    if not request.user.is_staff:
        messages.error(
            request, "Acceso denegado: solo los administradores pueden eliminar instalaciones.")
        return redirect('lista_instalaciones')

    instalacion = get_object_or_404(CuadroInsta, id=id)
    instalacion.delete()
    messages.success(request, "Instalación eliminada correctamente.")
    return redirect('lista_instalaciones')


# se utiliza para importar los archivos excel en la vista de carga de excel 2 de mayo 2025


# ------------------ Utilidades ------------------
DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"]
DATETIME_FORMATS = ["%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S",
                    "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M"]


def _is_empty(v):
    return v is None or (isinstance(v, float) and str(v) == 'nan') or (isinstance(v, str) and v.strip() == "")


def to_date(value):
    """Normaliza a datetime.date (None si no parsea)."""
    if _is_empty(value):
        return None
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        s = value.strip()
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    try:
        return datetime.fromisoformat(str(value)).date()
    except Exception:
        return None


def to_datetime(value, tz_aware=True):
    """Normaliza a datetime (timezone aware)."""
    if _is_empty(value):
        return None
    if isinstance(value, datetime):
        dt = value
    elif isinstance(value, date):
        dt = datetime.combine(value, time(0, 0))
    elif isinstance(value, str):
        s = value.strip()
        # ISO 8601 rápido
        try:
            dt = datetime.fromisoformat(s.replace("T", " "))
        except Exception:
            dt = None
        if dt is None:
            for fmt in DATETIME_FORMATS:
                try:
                    dt = datetime.strptime(s, fmt)
                    break
                except ValueError:
                    continue
        if dt is None:
            d = to_date(s)
            if d:
                dt = datetime.combine(d, time(0, 0))
            else:
                return None
    else:
        return None

    if tz_aware and timezone.is_naive(dt):
        dt = timezone.make_aware(dt, timezone.get_current_timezone())
    return dt


def norm_text(s):
    return None if _is_empty(s) else str(s).strip()


def norm_si_no(v):
    if _is_empty(v):
        return "NO"
    s = str(v).strip().upper()
    return "SI" if s in ("SI", "SÍ", "YES", "Y", "1", "TRUE") else "NO"


def find_tecnico_by_value(v, fallback_map=None):
    """
    Acepta: ID numérico, nombre (case-insensitive), o mapea por diccionario.
    """
    if _is_empty(v):
        return None
    if isinstance(v, (int,)) or (isinstance(v, str) and v.isdigit()):
        pk = int(v)
        return Tecnico.objects.filter(pk=pk).first()  # pylint: disable=no-member
    if fallback_map:
        m = fallback_map.get(str(v).strip(), None)
        if m:
            return Tecnico.objects.filter(pk=m).first()  # pylint: disable=no-member

    return Tecnico.objects.filter(nombre__iexact=str(v).strip()).first()  # pylint: disable=no-member


def find_ejecutivo_by_value(v, fallback_map=None):
    if _is_empty(v):
        return None
    if isinstance(v, (int,)) or (isinstance(v, str) and v.isdigit()):
        pk = int(v)
        return Ejecutivo.objects.filter(pk=pk).first()  # pylint: disable=no-member
    if fallback_map:
        m = fallback_map.get(str(v).strip(), None)
        if m:
            return Ejecutivo.objects.filter(pk=m).first()  # pylint: disable=no-member
    return Ejecutivo.objects.filter(nombre__iexact=str(v).strip()).first()  # pylint: disable=no-member


def normalize_header(h):
    return str(h).strip().lower() if h is not None else ""


# ------------------ Vista de importación ------------------

@user_passes_test(lambda u: u.is_superuser)
def importar_excel(request):
    if request.method == 'POST':
        form = ExcelUploadForm(request.POST, request.FILES)
        if form.is_valid():
            archivo = request.FILES['archivo_excel']
            nombre = archivo.name

            # Evitar duplicar el mismo archivo
            if RegistroImportacion.objects.filter(nombre_archivo=nombre).exists():  # pylint: disable=no-member
                messages.warning(
                    request, f"El archivo '{nombre}' ya fue cargado anteriormente.")
                return redirect('lista_instalaciones')

            try:
                wb = openpyxl.load_workbook(archivo, data_only=True)
            except Exception as e:
                messages.error(request, f"No se pudo abrir el Excel: {e}")
                return redirect('lista_instalaciones')

            hoja = wb.active

            # Leer encabezados
            headers_row = list(hoja.iter_rows(
                min_row=1, max_row=1, values_only=True))[0]
            header_map = {normalize_header(
                h): idx for idx, h in enumerate(headers_row)}

            def get(row, *names):
                for n in names:
                    idx = header_map.get(normalize_header(n))
                    if idx is not None and idx < len(row):
                        return row[idx]
                return None

            # Fallbacks opcionales (si aún los quieres)
            tecnico_map = {
                "Ricardo": 1, "Giovanny": 2, "Juan": 3, "Francisco": 4,
                "German": 5, "Willinthon": 6, "Jhon": 7, "Guido": 8,
                "Esneyder": 9, "Diego": 10, "-----": 11
            }
            ejecutivo_map = {
                "Francisco Silva": 1, "Andrea Trujillo": 2,
                "Yamile David": 3, "-----": 4
            }

            errores = []

            for idx, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
                if not any(fila):
                    continue  # fila vacía

                try:
                    # Lee por nombre (plantilla v2) y soporta plantilla antigua
                    pvg = get(fila, "PVG")
                    # fecha datetime
                    ingreso = get(fila, "Ingreso", "Fecha", "fecha")
                    codigo = get(fila, "Código", "Codigo", "codigo", "CODIGO")
                    cliente = get(fila, "Cliente", "cliente", "CLIENTE")
                    ciudad = get(fila, "Ciudad", "ciudad", "CIUDAD")
                    direccion = get(fila, "Dirección",
                                    "Direccion", "direccion", "DIRECCION")
                    instalacion = get(fila, "Instalación",
                                      "Instalacion", "instalacion")
                    dias = get(fila, "Días", "Dias", "dias_cotizados")

                    bodega = get(fila, "Bodega", "en_bodega")
                    inicio = get(fila, "Inicio", "fecha_inicio")
                    orden = get(fila, "Orden", "orden")
                    obs = get(fila, "Observación",
                              "Observacion", "observacion")
                    ejecutivo_v = get(fila, "Ejecutivo", "ejecutivo")

                    tec1_v = get(fila, "Técnico 1")
                    tec2_v = get(fila, "Técnico 2")
                    if tec1_v is None and tec2_v is None:
                        tecnicos_comb = get(fila, "Técnicos")
                        if tecnicos_comb:
                            partes = [t.strip() for t in str(tecnicos_comb).replace(
                                "/", ",").split(",") if t.strip()]
                            tec1_v = partes[0] if len(partes) > 0 else None
                            tec2_v = partes[1] if len(partes) > 1 else None

                    ciudad_norm = norm_text(ciudad)
                    ciudad_obj = None

                    if ciudad_norm:
                        ciudad_obj = Ciudad.objects.filter(
                            nombre__iexact=ciudad_norm.strip()
                        ).first()

                        if not ciudad_obj:
                            ciudad_obj = Ciudad.objects.create(
                                nombre=ciudad_norm.strip().upper()
                            )

                    fecha_dt = to_datetime(ingreso)
                    fecha_ini = to_date(inicio)
                    en_bodega = norm_si_no(bodega)

                    # Pre-chequeo de duplicados por (PVG, Ciudad) para informar antes
                    if pvg not in (None, "") and ciudad_norm:
                        if CuadroInsta.objects.filter(pvg=pvg, ciudad__iexact=ciudad_norm).exists():  # pylint: disable=no-member
                            mensaje = f"Duplicado: ya existe PVG {pvg} en ciudad '{ciudad_norm}'."
                            errores.append((idx, fila, mensaje))
                            messages.warning(request, f"Fila {idx}: {mensaje}")
                            continue

                    # Resolver FKs (ID o nombre)
                    t1 = find_tecnico_by_value(tec1_v, tecnico_map)
                    t2 = find_tecnico_by_value(tec2_v, tecnico_map)
                    ej = find_ejecutivo_by_value(ejecutivo_v, ejecutivo_map)

                    data = {
                        "pvg": pvg if pvg not in (None, "") else None,
                        "fecha": fecha_dt,
                        "codigo": norm_text(codigo),
                        "cliente": norm_text(cliente),
                        "ciudad": ciudad_obj.pk if ciudad_obj else None,
                        "direccion": norm_text(direccion),
                        "instalacion": norm_text(instalacion),
                        "dias_cotizados": dias,

                        "en_bodega": en_bodega,
                        "fecha_inicio": fecha_ini,
                        "orden": norm_text(orden),
                        "observacion": norm_text(obs),
                        "tecnico1": t1.pk if t1 else None,
                        "tecnico2": t2.pk if t2 else None,
                        "ejecutivo": ej.pk if ej else None,
                    }

                    # Usa el ModelForm con modo_import para la política de fechas y reglas de negocio
                    form = CuadroInstaForm(data=data, modo_import=True)
                    if form.is_valid():
                        obj = form.save(commit=False)
                        obj.usuario = request.user
                        obj.save()
                    else:
                        detalles = []
                        for campo, msgs in form.errors.items():
                            for m in msgs:
                                detalles.append(f"{campo}: {m}")
                        raise ValueError("; ".join(detalles))

                except Exception as e:
                    errores.append((idx, fila, str(e)))
                    messages.error(request, f"Error en fila {idx}: {e}")

            # Registrar importación
            RegistroImportacion.objects.create(  # pylint: disable=no-member
                nombre_archivo=nombre,
                usuario=request.user
            )

            # Si hubo errores, devolver CSV
            if errores:
                response = HttpResponse(content_type='text/csv')
                response[
                    'Content-Disposition'] = f'attachment; filename=errores_importacion_{nombre}.csv'
                writer = csv.writer(response)
                writer.writerow(['Fila', 'Contenido', 'Error'])
                for fila_num, contenido, mensaje_error in errores:
                    writer.writerow([fila_num, list(contenido), mensaje_error])
                return response

            messages.success(
                request, f"El archivo '{nombre}' fue importado exitosamente.")
            return redirect('lista_instalaciones')

    else:
        form = ExcelUploadForm()

    return render(request, 'importar_excel.html', {'form': form})


def exportar_excel(request):
    instalaciones = CuadroInsta.objects.all().values(  # pylint: disable=no-member
        'pvg', 'fecha', 'codigo', 'cliente', 'ciudad', 'direccion',
        'instalacion', 'dias_cotizados', 'cantidad_tecnicos', 'en_bodega',
        'fecha_inicio', 'fecha_terminacion', 'finaliza', 'orden',
        'tecnico1', 'tecnico2', 'estado', 'observacion', 'ejecutivo'
    )

    df = pd.DataFrame(list(instalaciones))

   # Convertir columnas datetime a solo fecha (sin hora ni zona horaria)
    for col in ['fecha', 'fecha_inicio', 'fecha_terminacion', 'finaliza']:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors='coerce').dt.date

    response = HttpResponse(content_type='application/vnd.ms-excel')
    response['Content-Disposition'] = 'attachment; filename="instalaciones.xlsx"'
    df.to_excel(response, index=False)
    return response

# HOJA DE MANTENIMIENTOS 2 AGOSTO


@login_required(login_url='login')
def lista_mantenimientos(request):

    contexto = {
        # luego aquí pondrás mantenimientos, tecnicos, etc.
    }

    return render(
        request,
        '/mantenimientos/listar_mantenimientos.html',
        contexto
    )
